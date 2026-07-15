"""Genera el Libro de Compras en .xlsx para un periodo (mes) dado.

El Libro de Compras del SII debe ir en pesos. Por eso, ademas de los montos en la
moneda original de la factura, se agregan columnas convertidas a CLP usando el
tipo de cambio (dolar observado) guardado por factura.

La conversion a CLP va como FORMULA (=ROUND(monto*TC,0)), para que quede a la
vista y sea auditable. El afecto CLP se deriva del total (=total-iva-exento) para
que neto+iva+exento cuadre exacto por fila. Se activa el recalculo al abrir
(fullCalcOnLoad) para que Excel muestre los valores. Nota: los visores que NO
calculan formulas (ej. la vista previa de Dropbox) mostraran esas celdas vacias;
hay que abrir el archivo en Excel/LibreOffice/Sheets para ver los pesos.
"""
import io
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.models import Purchase

HEADERS = [
    "Index", "NUMERO", "FECHA", "RUT", "DESCRIPCION", "MONEDA",
    "AFECTO", "EXENTO", "IVA", "TOTAL",              # montos en moneda original (G,H,I,J)
    "T/CAMBIO", "FECHA T/C",                          # dolar observado aplicado (K,L)
    "AFECTO CLP", "EXENTO CLP", "IVA CLP", "TOTAL CLP",  # convertidos a pesos (M,N,O,P)
]

FMT_ORIG = "#,##0.00"
FMT_CLP = "#,##0"
FMT_TC = "#,##0.00"


def build_workbook(purchases: Sequence[Purchase], periodo_label: str) -> bytes:
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True  # que Excel recalcule las formulas al abrir
    ws = wb.active
    ws.title = "Pag1"

    ws.merge_cells("E1:J1")
    title_cell = ws["E1"]
    title_cell.value = f"LIBRO COMPRA - {periodo_label}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")
    note = ws["E2"]
    note.value = "Montos originales por moneda; columnas 'CLP' convertidas al dolar observado (para el libro en pesos)."
    note.font = Font(italic=True, size=9)

    header_row = 3
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True)

    row = header_row + 1
    first_data_row = row
    for idx, p in enumerate(purchases, start=1):
        r = row
        es_usd = (p.moneda or "").upper() == "USD"
        tc = float(p.tipo_cambio) if p.tipo_cambio is not None else None

        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=p.numero)
        ws.cell(row=r, column=3, value=p.fecha.strftime("%d-%m-%Y") if p.fecha else None)
        ws.cell(row=r, column=4, value=p.rut_proveedor)
        descripcion = p.proveedor or ""
        if p.descripcion:
            descripcion = f"{descripcion} - {p.descripcion}"
        ws.cell(row=r, column=5, value=descripcion)
        ws.cell(row=r, column=6, value=p.moneda)

        # Montos en moneda original (G,H,I,J)
        for col, val in ((7, p.monto_afecto), (8, p.monto_exento), (9, p.iva), (10, p.total)):
            c = ws.cell(row=r, column=col, value=float(val) if val is not None else None)
            c.number_format = FMT_ORIG

        # Columnas en CLP como formula (N exento, O iva, P total; M afecto derivado)
        if es_usd and tc:
            tcc = ws.cell(row=r, column=11, value=tc)
            tcc.number_format = FMT_TC
            ws.cell(row=r, column=12,
                    value=p.tipo_cambio_fecha.strftime("%d-%m-%Y") if p.tipo_cambio_fecha else None)
            ws.cell(row=r, column=14, value=f"=ROUND(H{r}*$K{r},0)")  # exento
            ws.cell(row=r, column=15, value=f"=ROUND(I{r}*$K{r},0)")  # iva
            ws.cell(row=r, column=16, value=f"=ROUND(J{r}*$K{r},0)")  # total
            ws.cell(row=r, column=13, value=f"=P{r}-O{r}-N{r}")       # afecto = total - iva - exento
            clp_cols = (13, 14, 15, 16)
        elif not es_usd:
            # Ya esta en CLP: la columna convertida referencia el monto original
            ws.cell(row=r, column=14, value=f"=H{r}")
            ws.cell(row=r, column=15, value=f"=I{r}")
            ws.cell(row=r, column=16, value=f"=J{r}")
            ws.cell(row=r, column=13, value=f"=P{r}-O{r}-N{r}")
            clp_cols = (13, 14, 15, 16)
        else:
            clp_cols = ()  # USD sin TC: no se puede convertir, columnas CLP en blanco

        for col in clp_cols:
            ws.cell(row=r, column=col).number_format = FMT_CLP
        row += 1

    last_data_row = row - 1
    total_row = row + 1
    ws.cell(row=total_row, column=2, value="Total :")
    ws.cell(row=total_row, column=4, value=f"{len(purchases)} Documentos")
    ws.cell(row=total_row, column=5, value="Total General (CLP)")
    if last_data_row >= first_data_row:
        for col in (13, 14, 15, 16):  # solo columnas CLP: el libro suma en pesos
            letter = get_column_letter(col)
            c = ws.cell(row=total_row, column=col,
                        value=f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
            c.number_format = FMT_CLP
            c.font = Font(bold=True)

    widths = [7, 16, 12, 14, 40, 8, 12, 12, 12, 12, 10, 12, 14, 12, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
